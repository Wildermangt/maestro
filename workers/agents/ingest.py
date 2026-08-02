"""
Agente Ingesta — lee archivos aportados por el usuario.

Hasta ahora el sistema solo podía consumir la web. Con esto puede partir
de tus propios documentos: una lista de precios en Excel, una ficha
técnica en PDF, un CSV de clientes.

Formatos: PDF, DOCX, XLSX, CSV y texto plano.

Seguridad — dos reglas, y las dos importan:

1. **Solo se lee dentro de `INGEST_DIR`.** Toda ruta se resuelve y se
   comprueba que quede dentro de esa carpeta. Sin esto, un prompt como
   "lee /app/.env" o "lee ../../etc/passwd" haría que el agente
   volcara secretos al contexto del LLM, que luego viajan al proveedor.

2. **No se ejecuta nada de lo que se lee.** Solo extracción de texto.

El agente NO recibe rutas del LLM: recibe los nombres de archivo en
`metadata.files`, que los pone quien crea la tarea. El prompt del
usuario solo sirve para decidir qué hacer con el contenido.
"""
import csv
import io
import logging
import os
from pathlib import Path

from agents.base import BaseAgent
from models import IngestOutput, TaskJob, TaskResult, TaskStatus

logger = logging.getLogger("ingest")

INGEST_ROOT = Path(os.getenv("INGEST_DIR", "/app/ingest")).resolve()

# Tope por archivo y total, para no reventar el contexto del LLM que
# recibirá esto después.
MAX_CHARS_POR_ARCHIVO = 40000
MAX_CHARS_TOTAL = 120000

EXTENSIONES = {".pdf", ".docx", ".xlsx", ".xls", ".csv", ".txt", ".md", ".json"}


class IngestAgent(BaseAgent):
    name = "ingest"

    def run(self, job: TaskJob) -> TaskResult:
        nombres = job.metadata.get("files") or []
        logger.info(f"IngestAgent procesando taskId={job.taskId} archivos={nombres}")

        if not nombres:
            disponibles = self._listar_disponibles()
            return TaskResult(
                taskId=job.taskId,
                status=TaskStatus.FAILED,
                result={"archivosDisponibles": disponibles},
                error=(
                    "No se indicó ningún archivo. Pasa los nombres en metadata.files. "
                    f"Archivos disponibles en la carpeta de ingesta: {disponibles or 'ninguno'}."
                ),
            )

        piezas: list[str] = []
        leidos: list[str] = []
        errores: list[str] = []
        total = 0

        for nombre in nombres:
            try:
                ruta = self._ruta_segura(nombre)
            except ValueError as exc:
                errores.append(str(exc))
                continue

            try:
                texto = self._leer(ruta)
            except Exception as exc:
                logger.exception(f"No se pudo leer {ruta.name}")
                errores.append(f"{ruta.name}: {exc}")
                continue

            texto = texto[:MAX_CHARS_POR_ARCHIVO]
            if total + len(texto) > MAX_CHARS_TOTAL:
                texto = texto[: max(0, MAX_CHARS_TOTAL - total)]
                errores.append(f"{ruta.name}: recortado, se alcanzó el tope total de contenido.")

            total += len(texto)
            leidos.append(ruta.name)
            piezas.append(f"=== ARCHIVO: {ruta.name} ===\n{texto}")

            if total >= MAX_CHARS_TOTAL:
                break

        if not leidos:
            return TaskResult(
                taskId=job.taskId,
                status=TaskStatus.FAILED,
                result={},
                error="No se pudo leer ningún archivo. " + " | ".join(errores),
            )

        salida = IngestOutput(
            summary=f"Se leyeron {len(leidos)} archivo(s): {', '.join(leidos)}.",
            files_read=leidos,
            content="\n\n".join(piezas),
        )
        return TaskResult(
            taskId=job.taskId,
            status=TaskStatus.PARTIAL if errores else TaskStatus.COMPLETED,
            result=salida.model_dump(),
            error=" | ".join(errores) if errores else None,
        )

    # --- Seguridad ----------------------------------------------------------

    @staticmethod
    def _ruta_segura(nombre: str) -> Path:
        """
        Resuelve el nombre dentro de INGEST_ROOT y verifica que no se salga.
        Se descarta cualquier componente de directorio antes de resolver,
        así que "../../etc/passwd" se convierte en "passwd" y se busca
        dentro de la carpeta de ingesta, donde no existirá.
        """
        base = os.path.basename(str(nombre).strip())
        if not base:
            raise ValueError(f"Nombre de archivo vacío: {nombre!r}")

        ruta = (INGEST_ROOT / base).resolve()
        if not str(ruta).startswith(str(INGEST_ROOT) + os.sep):
            raise ValueError(f"Ruta fuera de la carpeta de ingesta: {nombre!r}")
        if ruta.suffix.lower() not in EXTENSIONES:
            raise ValueError(f"Extensión no admitida: {ruta.suffix or 'sin extensión'}")
        if not ruta.is_file():
            raise ValueError(f"No existe el archivo: {base}")
        return ruta

    @staticmethod
    def _listar_disponibles() -> list[str]:
        if not INGEST_ROOT.is_dir():
            return []
        return sorted(
            p.name for p in INGEST_ROOT.iterdir()
            if p.is_file() and p.suffix.lower() in EXTENSIONES
        )

    # --- Lectura por formato -------------------------------------------------

    @classmethod
    def _leer(cls, ruta: Path) -> str:
        ext = ruta.suffix.lower()
        if ext == ".pdf":
            return cls._leer_pdf(ruta)
        if ext == ".docx":
            return cls._leer_docx(ruta)
        if ext in (".xlsx", ".xls"):
            return cls._leer_xlsx(ruta)
        if ext == ".csv":
            return cls._leer_csv(ruta)
        return ruta.read_text(encoding="utf-8", errors="replace")

    @staticmethod
    def _leer_pdf(ruta: Path) -> str:
        from pypdf import PdfReader

        lector = PdfReader(str(ruta))
        return "\n".join((p.extract_text() or "") for p in lector.pages)

    @staticmethod
    def _leer_docx(ruta: Path) -> str:
        from docx import Document

        doc = Document(str(ruta))
        partes = [p.text for p in doc.paragraphs if p.text.strip()]
        for tabla in doc.tables:
            for fila in tabla.rows:
                partes.append(" | ".join(c.text.strip() for c in fila.cells))
        return "\n".join(partes)

    @staticmethod
    def _leer_xlsx(ruta: Path) -> str:
        from openpyxl import load_workbook

        wb = load_workbook(str(ruta), read_only=True, data_only=True)
        partes = []
        for ws in wb.worksheets:
            partes.append(f"--- Hoja: {ws.title} ---")
            for fila in ws.iter_rows(values_only=True):
                if any(v is not None and str(v).strip() for v in fila):
                    partes.append(" | ".join("" if v is None else str(v) for v in fila))
        wb.close()
        return "\n".join(partes)

    @staticmethod
    def _leer_csv(ruta: Path) -> str:
        texto = ruta.read_text(encoding="utf-8-sig", errors="replace")
        try:
            dialecto = csv.Sniffer().sniff(texto[:4000])
        except csv.Error:
            dialecto = csv.excel
        filas = list(csv.reader(io.StringIO(texto), dialecto))
        return "\n".join(" | ".join(f) for f in filas)
