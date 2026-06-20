"""
Agente Presentador (Sprint 4).

Flujo:
1. Claude genera una estructura de slides (PresentationPlan: título +
   lista de slides, cada uno con título y bullets) a partir del prompt
   (que normalmente es el resumen de una investigación previa, si viene
   delegado por el Director — ver metadata.context_data).
2. python-pptx construye el archivo .pptx real, con un diseño simple
   pero limpio (título + bullets, una slide de portada).
3. El archivo se guarda en el volumen compartido (utils/artifact_storage)
   y se referencia en TaskResult.artifacts para que NestJS cree el
   registro Artifact descargable.

Diseño deliberadamente simple para el Sprint 4: sin imágenes generadas,
sin gráficos (Plotly/charts es una extensión natural si se necesita
más adelante), un solo layout de slide. El riesgo técnico real que
este agente prueba es "LLM genera estructura -> librería de Office
construye el archivo real -> el archivo es descargable", no la
elaboración visual.
"""
import json
import logging

from pptx import Presentation
from pptx.util import Inches, Pt

from agents.base import BaseAgent
from models import PresentationOutput, PresentationPlan, TaskJob, TaskResult, TaskStatus
from utils.artifact_storage import save_artifact
from utils.llm_factory import get_llm_client

logger = logging.getLogger("presenter")

SYSTEM_PROMPT = """Eres un Presentador experto en storytelling y síntesis \
ejecutiva. Recibes un objetivo o un resumen de investigación, y debes \
estructurarlo como una presentación clara y persuasiva.

Genera entre 4 y 8 slides (sin contar la portada). Cada slide debe tener \
un título corto y entre 2 y 5 bullets concisos (máximo ~15 palabras cada \
uno, sin párrafos largos — esto es una slide, no un documento).

Responde ÚNICAMENTE con un JSON con este shape exacto, sin texto adicional, \
sin backticks de markdown:

{
  "deck_title": "string, título general de la presentación",
  "slides": [
    {"title": "string", "bullets": ["string", "string", ...]}
  ]
}"""


class PresenterAgent(BaseAgent):
    name = "presenter"

    def __init__(self):
        self.llm = get_llm_client(provider="claude")

    def run(self, job: TaskJob) -> TaskResult:
        prompt = job.prompt
        context_data = job.metadata.get("context_data")
        logger.info(f"PresenterAgent procesando taskId={job.taskId} prompt='{prompt}'")

        user_message = f"OBJETIVO/TEMA: {prompt}"
        if context_data:
            user_message += f"\n\nCONTENIDO DE BASE (de investigación previa):\n{json.dumps(context_data, ensure_ascii=False)}"

        try:
            plan = self._generate_plan(user_message)
        except Exception as exc:
            logger.exception(f"No se pudo generar el plan de slides para taskId={job.taskId}")
            return TaskResult(
                taskId=job.taskId,
                status=TaskStatus.FAILED,
                result={},
                error=f"No se pudo generar la estructura de la presentación: {exc}",
            )

        try:
            pptx_bytes = self._build_pptx(plan)
        except Exception as exc:
            logger.exception(f"python-pptx falló construyendo el archivo para taskId={job.taskId}")
            return TaskResult(
                taskId=job.taskId,
                status=TaskStatus.FAILED,
                result={"plan": plan.model_dump()},
                error=f"La estructura se generó pero el archivo .pptx falló al construirse: {exc}",
            )

        artifact_ref = save_artifact(
            task_id=job.taskId,
            filename="presentacion.pptx",
            content_bytes=pptx_bytes,
            artifact_type="pptx",
        )

        output = PresentationOutput(
            summary=f"Presentación '{plan.deck_title}' generada con {len(plan.slides)} slide(s).",
            slide_count=len(plan.slides),
            artifactId="",  # NestJS asigna el ID real al crear el Artifact en Postgres
        )

        return TaskResult(
            taskId=job.taskId,
            status=TaskStatus.COMPLETED,
            result=output.model_dump(),
            artifacts=[artifact_ref.to_dict()],
        )

    def _generate_plan(self, user_message: str) -> PresentationPlan:
        raw = self.llm.complete(system=SYSTEM_PROMPT, user=user_message, max_tokens=2000)
        cleaned = self._strip_markdown_fence(raw)
        parsed = json.loads(cleaned)
        return PresentationPlan.model_validate(parsed)

    def _build_pptx(self, plan: PresentationPlan) -> bytes:
        import io

        prs = Presentation()

        # Slide de portada
        title_slide_layout = prs.slide_layouts[0]
        slide = prs.slides.add_slide(title_slide_layout)
        slide.shapes.title.text = plan.deck_title
        if len(slide.placeholders) > 1:
            slide.placeholders[1].text = "Generado por Prompt Maestro"

        # Slides de contenido
        bullet_layout = prs.slide_layouts[1]  # "Title and Content"
        for slide_content in plan.slides:
            slide = prs.slides.add_slide(bullet_layout)
            slide.shapes.title.text = slide_content.title

            body = slide.placeholders[1]
            text_frame = body.text_frame
            text_frame.clear()

            for i, bullet in enumerate(slide_content.bullets):
                p = text_frame.paragraphs[0] if i == 0 else text_frame.add_paragraph()
                p.text = bullet
                p.font.size = Pt(18)

        buffer = io.BytesIO()
        prs.save(buffer)
        return buffer.getvalue()

    @staticmethod
    def _strip_markdown_fence(text: str) -> str:
        cleaned = text.strip()
        if cleaned.startswith("```"):
            lines = cleaned.split("\n")
            lines = lines[1:] if lines[0].startswith("```") else lines
            if lines and lines[-1].strip() == "```":
                lines = lines[:-1]
            cleaned = "\n".join(lines)
        return cleaned
